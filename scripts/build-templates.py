"""Generate the OMT WORKING GROUP downloadable templates in assets/downloads/.

Run from the repository root:
    pip install openpyxl python-docx python-pptx
    python3 scripts/build-templates.py
"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor, Cm
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from pptx import Presentation
from pptx.dml.color import RGBColor as PptColor
from pptx.util import Inches, Pt as PptPt

OUT = Path(__file__).resolve().parent.parent / "assets" / "downloads"
OUT.mkdir(parents=True, exist_ok=True)

NAVY = "0F2A4A"
BLUE = "2563EB"
LIGHT = "EEF3FB"
CREDIT = "Template by OMT WORKING GROUP — free to use and adapt (CC BY 4.0)."


# ---------- Excel helpers ----------
def header_row(ws, row, values):
    thin = Side(style="thin", color="C9D3E3")
    for col, value in enumerate(values, start=1):
        cell = ws.cell(row=row, column=col, value=value)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(vertical="center")
        cell.border = Border(bottom=thin)


def title(ws, text, subtitle):
    ws["A1"] = text
    ws["A1"].font = Font(bold=True, size=16, color=NAVY)
    ws["A2"] = subtitle
    ws["A2"].font = Font(italic=True, color="5B6577")


def build_budget():
    wb = Workbook()
    ws = wb.active
    ws.title = "Budget"
    title(ws, "Student Monthly Budget", "Type in the yellow cells. Totals update automatically.")
    input_fill = PatternFill("solid", fgColor="FFF7D6")
    money = '#,##0.00'

    header_row(ws, 4, ["Income", "Planned", "Actual", "Difference"])
    incomes = ["Allowance / family support", "Part-time job", "Scholarship / bursary", "Other income"]
    for i, name in enumerate(incomes, start=5):
        ws.cell(row=i, column=1, value=name)
        for col in (2, 3):
            c = ws.cell(row=i, column=col, value=0)
            c.fill = input_fill
            c.number_format = money
        ws.cell(row=i, column=4, value=f"=C{i}-B{i}").number_format = money
    total_income_row = 5 + len(incomes)
    ws.cell(row=total_income_row, column=1, value="Total income").font = Font(bold=True)
    for col, letter in ((2, "B"), (3, "C"), (4, "D")):
        c = ws.cell(row=total_income_row, column=col, value=f"=SUM({letter}5:{letter}{total_income_row - 1})")
        c.font = Font(bold=True)
        c.number_format = money

    start = total_income_row + 2
    header_row(ws, start, ["Spending category", "Planned", "Actual (from Transactions)", "Difference"])
    categories = ["Rent / hostel", "Food", "Transport", "Books & supplies", "Phone & data",
                  "Personal care", "Entertainment", "Savings", "Other"]
    first = start + 1
    for i, name in enumerate(categories, start=first):
        ws.cell(row=i, column=1, value=name)
        c = ws.cell(row=i, column=2, value=0)
        c.fill = input_fill
        c.number_format = money
        ws.cell(row=i, column=3, value=f"=SUMIF(Transactions!$C:$C,A{i},Transactions!$D:$D)").number_format = money
        ws.cell(row=i, column=4, value=f"=B{i}-C{i}").number_format = money
    last = first + len(categories) - 1
    total_row = last + 1
    ws.cell(row=total_row, column=1, value="Total spending").font = Font(bold=True)
    for col, letter in ((2, "B"), (3, "C"), (4, "D")):
        c = ws.cell(row=total_row, column=col, value=f"=SUM({letter}{first}:{letter}{last})")
        c.font = Font(bold=True)
        c.number_format = money

    bal = total_row + 2
    ws.cell(row=bal, column=1, value="Remaining balance (actual income − actual spending)").font = Font(bold=True, color=NAVY)
    c = ws.cell(row=bal, column=3, value=f"=C{total_income_row}-C{total_row}")
    c.font = Font(bold=True, size=13, color=NAVY)
    c.number_format = money
    ws.cell(row=bal + 2, column=1, value=CREDIT).font = Font(italic=True, size=9, color="5B6577")
    ws.column_dimensions["A"].width = 46
    for col in "BCD":
        ws.column_dimensions[col].width = 24

    tx = wb.create_sheet("Transactions")
    title(tx, "Transactions", "Log each purchase. The Category must match a name on the Budget sheet.")
    header_row(tx, 4, ["Date", "Description", "Category", "Amount"])
    dv = DataValidation(type="list", formula1=f"=Budget!$A${first}:$A${last}", allow_blank=True)
    tx.add_data_validation(dv)
    dv.add("C5:C500")
    examples = [("2026-10-01", "Hostel fee (example — delete me)", "Rent / hostel", 0),
                ("2026-10-02", "Groceries (example — delete me)", "Food", 0)]
    for r, (d, desc, cat, amt) in enumerate(examples, start=5):
        tx.cell(row=r, column=1, value=d)
        tx.cell(row=r, column=2, value=desc)
        tx.cell(row=r, column=3, value=cat)
        tx.cell(row=r, column=4, value=amt).number_format = money
    tx.column_dimensions["A"].width = 14
    tx.column_dimensions["B"].width = 40
    tx.column_dimensions["C"].width = 22
    tx.column_dimensions["D"].width = 14
    tx.freeze_panes = "A5"
    wb.save(OUT / "omt-student-budget.xlsx")


def build_planner():
    wb = Workbook()
    ws = wb.active
    ws.title = "Assignments"
    title(ws, "Assignments and Deadlines", "Days left uses TODAY(), so it updates each time you open the file.")
    header_row(ws, 4, ["Course", "Task", "Due date", "Status", "Days left"])
    dv = DataValidation(type="list", formula1='"Not started,In progress,Done"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("D5:D200")
    for r in range(5, 41):
        ws.cell(row=r, column=3).number_format = "yyyy-mm-dd"
        ws.cell(row=r, column=5, value=f'=IF(OR(C{r}="",D{r}="Done"),"",C{r}-TODAY())')
    ws["A5"], ws["B5"], ws["D5"] = "Example: ECO 201", "Essay 1 (delete this row)", "Not started"
    for col, width in zip("ABCDE", (22, 40, 14, 16, 12)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A5"

    g = wb.create_sheet("Grades")
    title(g, "Grade Tracker", "Enter scores out of 100 and weights as percentages (e.g. 30 for 30%).")
    header_row(g, 4, ["Course", "Assessment", "Score (0–100)", "Weight (%)"])
    g["A5"], g["B5"], g["D5"] = "Example: ECO 201", "Mid-semester test", 30
    g["A6"], g["B6"], g["D6"] = "Example: ECO 201", "Final exam", 70
    g["F4"], g["G4"] = "Course", "Weighted average"
    for cell in (g["F4"], g["G4"]):
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=BLUE)
    for r in range(5, 13):
        g.cell(row=r, column=7, value=(
            f'=IF(F{r}="","",IFERROR(SUMPRODUCT(($A$5:$A$200=F{r})*($C$5:$C$200<>"")*$C$5:$C$200*$D$5:$D$200)'
            f'/SUMPRODUCT(($A$5:$A$200=F{r})*($C$5:$C$200<>"")*$D$5:$D$200),"No scores yet"))'
        )).number_format = "0.0"
    g["F5"] = "Example: ECO 201"
    g["F14"] = "Type each course name once in column F to see its weighted average."
    g["F14"].font = Font(italic=True, size=9, color="5B6577")
    g["A40"] = CREDIT
    g["A40"].font = Font(italic=True, size=9, color="5B6577")
    for col, width in zip("ABCDEFG", (22, 28, 14, 12, 3, 22, 18)):
        g.column_dimensions[col].width = width
    wb.save(OUT / "omt-study-planner.xlsx")


# ---------- Word ----------
def style_doc(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    h1 = doc.styles["Heading 1"]
    h1.font.size = Pt(13)
    h1.font.color.rgb = RGBColor.from_string(NAVY)
    for section in doc.sections:
        section.left_margin = section.right_margin = Cm(2)
        section.top_margin = section.bottom_margin = Cm(1.8)


def build_cv():
    doc = Document()
    style_doc(doc)
    name = doc.add_paragraph()
    run = name.add_run("[Your Full Name]")
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor.from_string(NAVY)
    doc.add_paragraph("[City, Country]  •  [Phone]  •  [professional.email@example.com]  •  [LinkedIn or portfolio link]")

    sections = {
        "Profile": ["[Two or three sentences: what you study, your strongest skills, and the role or internship you are applying for.]"],
        "Education": ["[Programme, e.g. BSc Economics] — [Institution], [Start year]–[Expected end year]",
                      "• [Relevant result, e.g. current grade average or class of degree]",
                      "• [Relevant modules, final-year project or academic award]"],
        "Experience": ["[Role] — [Organisation], [Month Year]–[Month Year]",
                       "• [Action verb + what you did + result, e.g. Organised a workshop for 60 students]",
                       "• [Another achievement with a number if possible]"],
        "Leadership and Volunteering": ["[Role] — [Society, club or community group], [Dates]",
                                        "• [What you contributed]"],
        "Skills": ["Digital: [e.g. Excel (PivotTables, XLOOKUP), Word, PowerPoint]",
                   "Languages: [e.g. English (fluent), Yoruba (native), French (basic)]",
                   "Other: [e.g. public speaking, data analysis, graphic design]"],
        "References": ["Available on request."],
    }
    for heading, lines in sections.items():
        doc.add_heading(heading, level=1)
        for line in lines:
            doc.add_paragraph(line)
    foot = doc.sections[0].footer.paragraphs[0]
    foot.text = "Delete this footer before sending. " + CREDIT
    foot.runs[0].font.size = Pt(8)
    doc.save(OUT / "omt-student-cv.docx")


def build_letter():
    doc = Document()
    style_doc(doc)
    for line in ["[Your Name]", "[Your Address]", "[Phone]  •  [Email]"]:
        doc.add_paragraph(line).paragraph_format.space_after = Pt(0)
    doc.add_paragraph()
    doc.add_paragraph("[Date, e.g. 15 October 2026]")
    for line in ["[Recipient's Name]", "[Title / Position]", "[Organisation]", "[Address]"]:
        doc.add_paragraph(line).paragraph_format.space_after = Pt(0)
    doc.add_paragraph()
    subject = doc.add_paragraph()
    subject.add_run("Subject: [What the letter is about]").bold = True
    doc.add_paragraph("Dear [Title and Surname, or Sir or Madam],")
    doc.add_paragraph("[Opening paragraph: state why you are writing in one or two sentences.]")
    doc.add_paragraph("[Body: give the details the reader needs. Use one paragraph per topic and keep sentences clear and polite.]")
    doc.add_paragraph("[Closing paragraph: say what you would like to happen next and thank the reader for their time.]")
    doc.add_paragraph("Yours sincerely,")
    doc.add_paragraph("\n\n[Your Name]")
    foot = doc.sections[0].footer.paragraphs[0]
    foot.text = "Delete this footer before sending. " + CREDIT
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    foot.runs[0].font.size = Pt(8)
    doc.save(OUT / "omt-formal-letter.docx")


# ---------- PowerPoint ----------
def build_deck(filename, slides):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    navy = PptColor.from_string(NAVY)
    blue = PptColor.from_string(BLUE)

    for index, (heading, body, notes) in enumerate(slides):
        layout = prs.slide_layouts[0 if index == 0 else 1]
        slide = prs.slides.add_slide(layout)
        # accent bar
        bar = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(0.18))
        bar.fill.solid()
        bar.fill.fore_color.rgb = blue if index else navy
        bar.line.fill.background()
        slide.shapes.title.text = heading
        for p in slide.shapes.title.text_frame.paragraphs:
            for r in p.runs:
                r.font.color.rgb = navy
                r.font.bold = True
        placeholder = slide.placeholders[1]
        tf = placeholder.text_frame
        tf.clear()
        for i, line in enumerate(body):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.text = line
            para.font.size = PptPt(24 if index else 22)
        slide.notes_slide.notes_text_frame.text = notes + "\n\n" + CREDIT
    prs.save(OUT / filename)


def build_academic():
    build_deck("omt-academic-presentation.pptx", [
        ("[Presentation Title]", ["[Your name] • [Course code] • [Date]"], "Introduce yourself and the topic in one sentence."),
        ("Outline", ["Background", "Research question", "Method", "Findings", "Conclusion"], "Tell the audience the route you will take."),
        ("Background", ["[What the audience needs to know first]", "[Key term or context]"], "Keep this short — only what is needed to follow the talk."),
        ("Research Question", ["[One clear sentence stating the question or problem]"], "Say the question slowly. Everything else answers it."),
        ("Method", ["[Data or sources used]", "[How you analysed them]"], "Explain how you investigated, not every detail."),
        ("Finding 1", ["[Headline that states the result]", "[Insert a chart: Insert › Chart]"], "One finding per slide. Point to the part of the chart that matters."),
        ("Finding 2", ["[Headline that states the result]", "[Supporting evidence]"], "Connect this finding to the first."),
        ("Discussion", ["[What the results mean]", "[Limitations of your study]"], "Showing limitations builds credibility."),
        ("Conclusion", ["[Key point 1]", "[Key point 2]", "[Key point 3]"], "End with the three things you want people to remember."),
        ("References and Questions", ["[Author, Year. Title. Publisher.]", "Thank you — questions?"], "Use the citation style your course requires."),
    ])


def build_proposal():
    build_deck("omt-project-proposal.pptx", [
        ("[Project Name]", ["Project proposal • [Team or presenter names] • [Date]"], "Open with the project name and who is presenting."),
        ("The Problem", ["[Who is affected?]", "[Why does it matter? Use one number if you can]"], "Lead with the problem, not the solution."),
        ("Proposed Solution", ["[What you will build or do, in one sentence]", "[How it solves the problem]"], "Keep this concrete and simple."),
        ("Objectives", ["[Specific, measurable objective 1]", "[Objective 2]", "[Objective 3]"], "Use SMART objectives: specific, measurable, achievable, relevant, time-bound."),
        ("Scope", ["In scope: [ ]", "Out of scope: [ ]"], "Being clear about what you will not do prevents misunderstandings."),
        ("Timeline", ["Phase 1 — [ ] — [dates]", "Phase 2 — [ ] — [dates]", "Phase 3 — [ ] — [dates]"], "Consider a simple table or bar chart for the timeline."),
        ("Budget and Resources", ["[Item] — [cost]", "[Item] — [cost]", "Total: [ ]"], "Show the total clearly and be ready to justify each item."),
        ("Risks and Mitigation", ["Risk: [ ] → Plan: [ ]", "Risk: [ ] → Plan: [ ]"], "Panels trust teams that have thought about what could go wrong."),
        ("Our Ask", ["[The approval, funding or support you need]", "[Next step and date]"], "End with a clear request so the audience knows what decision you need."),
    ])


if __name__ == "__main__":
    build_budget()
    build_planner()
    build_cv()
    build_letter()
    build_academic()
    build_proposal()
    for f in sorted(OUT.iterdir()):
        print(f"{f.name}: {f.stat().st_size // 1024} KB")
